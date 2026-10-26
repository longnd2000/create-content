import re
import sys
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

RED_KEYWORD_COLOR = RGBColor(192, 0, 0) # #C00000 - Màu đỏ chuẩn thẩm định phân bổ SEO

def set_cell_background(cell, fill_hex):
    """Đặt màu nền cho ô trong bảng"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    """Đặt viền cho bảng"""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Đặt lề trong cho ô bảng (dxa)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def parse_inline_formatting(paragraph, text, kw_pattern=None, kw_set=None):
    """Xử lý in đậm (**text**), in nghiêng (*text*) và code (`text`). Key luôn giữ định dạng chữ đen thông thường."""
    def add_segment(seg, is_bold=False, is_italic=False, is_code=False):
        if not seg:
            return
        if is_code:
            run = paragraph.add_run(seg)
            run.font.name = 'Google Sans Flex'
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0, 0, 0)
            return

        if kw_pattern and kw_set:
            parts = re.split(kw_pattern, seg, flags=re.IGNORECASE)
            for part in parts:
                if not part:
                    continue
                run = paragraph.add_run(part)
                run.font.name = 'Google Sans Flex'
                run.font.size = Pt(13)
                run.font.color.rgb = RGBColor(0, 0, 0)
                if is_bold:
                    run.bold = True
                if is_italic:
                    run.italic = True
        else:
            run = paragraph.add_run(seg)
            run.font.name = 'Google Sans Flex'
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if is_bold:
                run.bold = True
            if is_italic:
                run.italic = True

    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            add_segment(token[2:-2], is_bold=True)
        elif token.startswith('*') and token.endswith('*'):
            add_segment(token[1:-1], is_italic=True)
        elif token.startswith('`') and token.endswith('`'):
            add_segment(token[1:-1], is_code=True)
        else:
            add_segment(token, is_bold=False, is_italic=False)

def convert_md_to_docx(md_path, docx_path, keywords=None):
    doc = Document()

    # Cài đặt Margin trang
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Đặt font chữ mặc định là Google Sans Flex cho toàn tài liệu
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Google Sans Flex'
    normal_font.size = Pt(13)
    normal_font.color.rgb = RGBColor(0, 0, 0)
    normal_style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Google Sans Flex')
    normal_style.paragraph_format.line_spacing = 1.25
    normal_style.paragraph_format.space_after = Pt(6)

    for list_style_name in ('List Bullet', 'List Number'):
        list_style = doc.styles[list_style_name]
        list_style.font.name = 'Google Sans Flex'
        list_style.font.size = Pt(13)
        list_style.font.color.rgb = RGBColor(0, 0, 0)
        list_style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Google Sans Flex')

    # Đọc file markdown
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    # Tự động trích xuất danh sách từ khóa mục tiêu nếu chưa truyền
    target_keywords = []
    if keywords:
        target_keywords = list(keywords)
    else:
        in_kv = False
        for line in lines:
            s = line.strip()
            if s.startswith('|') and 'KEY' in s.upper() and 'VOLUME' in s.upper():
                in_kv = True
                continue
            if in_kv:
                if not s.startswith('|'):
                    in_kv = False
                else:
                    cols = [c.strip() for c in s[1:-1].split('|')]
                    if cols and not re.match(r'^:?-+:?$', cols[0]):
                        kw = cols[0].strip()
                        if kw and kw.upper() != 'KEY' and len(kw) > 3 and kw not in target_keywords:
                            target_keywords.append(kw)
            if s.lower().startswith('key chính:'):
                kc = s.split(':', 1)[1].strip()
                if kc and kc not in target_keywords:
                    target_keywords.append(kc)

    # Sắp xếp từ khóa theo độ dài giảm dần để ưu tiên cụm dài trước
    target_keywords.sort(key=len, reverse=True)
    kw_pattern = None
    kw_set = set()
    if target_keywords:
        kw_pattern = r'(' + '|'.join(re.escape(k) for k in target_keywords) + r')'
        kw_set = {k.lower() for k in target_keywords}

    def inline(p, text):
        parse_inline_formatting(p, text, kw_pattern, kw_set)

    in_table = False
    table_lines = []

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            in_table = False
            return
        
        # Phân tích các hàng của bảng
        rows_data = []
        for line in table_lines:
            line_str = line.strip()
            if line_str.startswith('|') and line_str.endswith('|'):
                line_str = line_str[1:-1]
            cols = [c.strip() for c in line_str.split('|')]
            # Bỏ qua dòng phân cách | :--- | :--- |
            if all(re.match(r'^:?-+:?$', c) for c in cols):
                continue
            rows_data.append(cols)

        if rows_data:
            max_cols = max(len(r) for r in rows_data)
            table = doc.add_table(rows=len(rows_data), cols=max_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True

            # Kiểm tra xem đây có phải bảng KEY & VOLUME không
            is_key_vol_table = False
            if rows_data and len(rows_data[0]) >= 2:
                header_texts = [c.strip().upper() for c in rows_data[0]]
                if any('KEY' in h for h in header_texts) and any('VOLUME' in h for h in header_texts):
                    is_key_vol_table = True

            set_table_borders(table, color="000000", sz="4", val="single")

            for row_idx, row_cells in enumerate(rows_data):
                is_header = (row_idx == 0)
                for col_idx in range(max_cols):
                    cell_text = row_cells[col_idx] if col_idx < len(row_cells) else ""
                    cell = table.cell(row_idx, col_idx)
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

                    if is_header:
                        set_cell_background(cell, "FFFFFF")
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run = p.add_run(cell_text)
                        run.bold = True
                        run.font.name = 'Google Sans Flex'
                        run.font.color.rgb = RGBColor(0, 0, 0)
                        run.font.size = Pt(13)
                    else:
                        if is_key_vol_table:
                            set_cell_background(cell, "FFFFFF")
                            if col_idx == 1:
                                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            else:
                                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        else:
                            set_cell_background(cell, "FFFFFF")
                        inline(p, cell_text)
                        for r in p.runs:
                            r.font.size = Pt(13)

            # Thêm khoảng cách sau bảng
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(6)

        table_lines = []
        in_table = False

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        stripped = line.strip()

        # Kiểm tra bảng markdown
        if stripped.startswith('|') and '|' in stripped[1:]:
            in_table = True
            table_lines.append(stripped)
            i += 1
            continue
        elif in_table:
            flush_table()

        if not stripped:
            i += 1
            continue

        # Thẻ Heading H1 - H6 hoặc dòng H1:
        h_match = re.match(r'^(#{1,6})\s+(.*)$', stripped)
        h1_line_match = re.match(r'^H1:\s*(.*)$', stripped, flags=re.IGNORECASE)
        if h_match or h1_line_match:
            if h_match:
                level = len(h_match.group(1))
                heading_text = h_match.group(2).strip()
            else:
                level = 1
                heading_text = stripped
            
            # Sử dụng chuẩn style Heading của Word: 'Heading 1', 'Heading 2', ...
            style_name = f'Heading {level}'
            try:
                p = doc.add_paragraph(style=style_name)
            except Exception:
                p = doc.add_paragraph()
            
            inline(p, heading_text)
            
            # Custom font, cỡ chữ và màu đen cho Heading
            for r in p.runs:
                r.bold = True
                if level == 1:
                    r.font.size = Pt(16)
                    r.font.color.rgb = RGBColor(0, 0, 0)
                    r.font.name = 'Google Sans Flex'
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(14)
                    p.paragraph_format.space_after = Pt(8)
                elif level == 2:
                    r.font.size = Pt(13)
                    r.font.color.rgb = RGBColor(0, 0, 0)
                    r.font.name = 'Google Sans Flex'
                    p.paragraph_format.space_before = Pt(12)
                    p.paragraph_format.space_after = Pt(6)
                elif level == 3:
                    r.font.size = Pt(13)
                    r.font.color.rgb = RGBColor(0, 0, 0)
                    r.font.name = 'Google Sans Flex'
                    p.paragraph_format.space_before = Pt(8)
                    p.paragraph_format.space_after = Pt(4)
                elif level == 4:
                    r.font.size = Pt(13)
                    r.font.color.rgb = RGBColor(0, 0, 0)
                    r.font.name = 'Google Sans Flex'
                    p.paragraph_format.space_before = Pt(6)
                    p.paragraph_format.space_after = Pt(3)
            i += 1
            continue

        # Xử lý Danh sách không thứ tự (Bullet points: -, *)
        bullet_match = re.match(r'^[-*]\s+(.*)$', stripped)
        if bullet_match:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            inline(p, bullet_match.group(1))
            i += 1
            continue

        # Xử lý Danh sách có thứ tự (Numbered list: 1., 2.)
        num_match = re.match(r'^\d+\.\s+(.*)$', stripped)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            inline(p, num_match.group(1))
            i += 1
            continue

        # Xử lý Blockquote (> text)
        if stripped.startswith('>'):
            quote_text = stripped[1:].strip()
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            set_cell_background(cell, "FFFFFF")
            set_cell_margins(cell, top=100, bottom=100, left=180, right=180)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run_prefix = p.add_run("LƯU Ý: ")
            run_prefix.bold = True
            run_prefix.font.color.rgb = RGBColor(0, 0, 0)
            inline(p, quote_text)
            
            # Thêm khoảng trống sau quote
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # Xử lý Khung Hướng dẫn Hình ảnh dạng cũ ([IMAGE_BRIEF] ... [/IMAGE_BRIEF])
        if stripped.startswith('[IMAGE_BRIEF]'):
            brief_lines = []
            i += 1
            while i < len(lines):
                b_line = lines[i].rstrip('\r\n').strip()
                if b_line.startswith('[/IMAGE_BRIEF]'):
                    i += 1
                    break
                if b_line:
                    brief_lines.append(b_line)
                i += 1
            
            if brief_lines:
                for b_text in brief_lines:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(6)
                    inline(p, b_text)
                    for r in p.runs:
                        r.font.name = 'Google Sans Flex'
                        r.font.size = Pt(13)
                        r.italic = True
                        r.font.color.rgb = RED_KEYWORD_COLOR
            continue

        # Xử lý ghi chú hình ảnh ngắn gọn dạng text thường trong ngoặc đơn, màu đỏ
        is_simple_image_note = (
            stripped.startswith('(') and stripped.endswith(')')
        )
        if is_simple_image_note or stripped.startswith('📸') or stripped.startswith('![') or ('Hình ảnh' in stripped and 'Alt text' in stripped):
            if stripped.startswith('📸'):
                stripped = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(6)
            inline(p, stripped)
            for r in p.runs:
                r.font.name = 'Google Sans Flex'
                r.font.size = Pt(13)
                r.italic = True
                r.font.color.rgb = RED_KEYWORD_COLOR
            i += 1
            continue

        # Đoạn văn bản thông thường (Paragraph)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        inline(p, stripped)
        i += 1

    # Nếu file kết thúc khi đang ở trong bảng
    if in_table:
        flush_table()

    out_dir = os.path.dirname(os.path.abspath(docx_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    doc.save(docx_path)
    print(f"Successfully generated DOCX: {docx_path}")

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python md_to_docx.py <input.md> <output.docx>")
        sys.exit(1)
    convert_md_to_docx(sys.argv[1], sys.argv[2])
