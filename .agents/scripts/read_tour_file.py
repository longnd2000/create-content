import sys
import os

def read_file(filepath):
    sys.stdout.reconfigure(encoding='utf-8')
    if not os.path.exists(filepath):
        print(f"Error: File '{filepath}' does not exist.")
        sys.exit(1)
    
    ext = os.path.splitext(filepath)[1].lower()
    
    if ext == '.docx':
        from docx import Document
        doc = Document(filepath)
        for p in doc.paragraphs:
            txt = p.text.strip()
            if txt:
                print(txt)
        for table in doc.tables:
            print("\n[TABLE]")
            for row in table.rows:
                row_txt = " | ".join(cell.text.strip().replace('\n', ' ') for cell in row.cells)
                print(f"| {row_txt} |")
            print("[/TABLE]\n")
    elif ext in ['.txt', '.md']:
        with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
            print(f.read())
    elif ext == '.pdf':
        try:
            import pypdf
            reader = pypdf.PdfReader(filepath)
            for page in reader.pages:
                print(page.extract_text())
        except Exception:
            try:
                import fitz
                doc = fitz.open(filepath)
                for page in doc:
                    print(page.get_text())
            except Exception as e:
                print(f"Cannot read PDF: {e}")
    else:
        with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
            print(f.read())

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python read_tour_file.py <path_to_file>")
        sys.exit(1)
    read_file(sys.argv[1])
