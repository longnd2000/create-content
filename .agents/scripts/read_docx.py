import sys
from docx import Document

def docx_to_text(docx_path):
    sys.stdout.reconfigure(encoding='utf-8')
    doc = Document(docx_path)
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        style = p.style.name.lower()
        if 'heading 1' in style:
            print(f"# {txt}\n")
        elif 'heading 2' in style:
            print(f"## {txt}\n")
        elif 'heading 3' in style:
            print(f"### {txt}\n")
        elif 'heading 4' in style:
            print(f"#### {txt}\n")
        elif 'bullet' in style:
            print(f"- {txt}")
        elif 'number' in style:
            print(f"1. {txt}")
        else:
            print(f"{txt}\n")
    
    # In nội dung các bảng
    for table in doc.tables:
        print("\n[TABLE]")
        for row in table.rows:
            row_txt = " | ".join(cell.text.strip().replace('\n', ' ') for cell in row.cells)
            print(f"| {row_txt} |")
        print("[/TABLE]\n")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python read_docx.py <path_to_docx>")
        sys.exit(1)
    docx_to_text(sys.argv[1])
